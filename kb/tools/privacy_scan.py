"""ONT.11 pseudonymization dry run on a case file, and the leak scan of generated outputs. Returns COUNTS ONLY, never the matched text.

1 closed dictionary from structured fields (names, every name token of two letters or more, plate); 2 pattern pass;
3 replace locally; 4 leak scan on the result with the same patterns (ONT.11.4), including name fragments right before or
right after a code. The mapping stays in memory and is never written. output_scan() runs the same checks over generated
files (workbook cells, reports), which must contain zero dictionary terms.
"""
from __future__ import annotations

import copy
import json
import re

PATTERNS = {
    "SG_PLATE": re.compile(r"\bS[A-Z]{1,2}\s?\d{1,4}\s?[A-Z]\b"),
    "PHONE": re.compile(r"(?:\+65[\s-]?)?\b[3689]\d{3}[\s-]?\d{4}\b"),
    "EMAIL": re.compile(r"\b[\w.+-]+@[\w-]+\.[\w.-]+\b"),
    "NRIC_FIN": re.compile(r"\b[STFGM]\d{7}[A-Z]\b"),
    "UNIT_NUMBER": re.compile(r"#\d{1,3}-\d{1,5}\b"),
    "LONG_NUMBER": re.compile(r"\b(?:\d[ -]?){12,19}\b"),
}
CODE = r"(?:RIDER|DRIVER)-NAME-\d+"
NOT_NAMES = {"Hi", "Hello", "Hey", "Dear", "Mr", "Mrs", "Ms", "Mdm", "Dr", "Sir", "Madam", "Uncle", "Auntie", "Thanks", "Thank",
             "Driver", "Rider", "Bro", "Ok", "Okay", "Please", "And", "The", "From", "To", "For", "With", "By"}


def _luhn(num: str) -> bool:
    digits = [int(c) for c in num if c.isdigit()]
    s = 0
    for i, d in enumerate(reversed(digits)):
        if i % 2:
            d = d * 2 - 9 if d * 2 > 9 else d * 2
        s += d
    return len(digits) >= 12 and s % 10 == 0


def dictionary(case: dict) -> list[tuple[str, str, str]]:
    """(term, code, type) from structured fields. Full names first, then tokens of two letters or more, then the plate."""
    terms = []
    for role, key in (("RIDER", "rider_profile"), ("DRIVER", "driver_profile")):
        name = case.get(key, {}).get("name")
        if name:
            terms.append((name, f"{role}-NAME-1", "PERSON_NAME"))
            for tok in re.split(r"[\s,.]+", name):
                if len(tok) >= 2:
                    terms.append((tok, f"{role}-NAME-1", "PERSON_NAME_TOKEN"))
    veh = case.get("driver_profile", {}).get("vehicle", "")
    m = re.search(r"\(([^)]+)\)", veh)
    if m:
        terms.append((m.group(1), "PLATE-1", "VEHICLE_PLATE"))
    return terms


def _strings(obj):
    if isinstance(obj, dict):
        for v in obj.values():
            yield from _strings(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from _strings(v)
    elif isinstance(obj, str):
        yield obj


def _count(strings, terms, allow_numbers=frozenset(), patterns=True) -> dict:
    counts = {"PERSON_NAME": 0, "PERSON_NAME_TOKEN": 0, "VEHICLE_PLATE": 0}
    counts.update({k: 0 for k in PATTERNS})
    for s in strings:
        for term, _, typ in terms:
            counts[typ] += len(re.findall(r"(?<![A-Za-z])" + re.escape(term) + r"(?![A-Za-z])", s, re.I))
        for k, p in (PATTERNS.items() if patterns else ()):
            hits = p.findall(s)
            if k == "LONG_NUMBER":
                hits = [h for h in hits if _luhn(h) and re.sub(r"\D", "", h) not in allow_numbers]
            counts[k] += len(hits)
    return counts


def scan(case: dict, terms) -> dict:
    return _count(list(_strings(case)), terms)


def pseudonymize(case: dict, terms) -> dict:
    out = copy.deepcopy(case)
    ordered = sorted(terms, key=lambda x: -len(x[0]))

    def repl(s: str) -> str:
        for term, code, _ in ordered:
            s = re.sub(r"(?<![A-Za-z])" + re.escape(term) + r"(?![A-Za-z])", code, s, flags=re.I)
        s = re.sub(r"\b(" + CODE + r")(\s+\1)+\b", r"\1", s)          # 'DRIVER-NAME-1 DRIVER-NAME-1' -> one code
        s = PATTERNS["SG_PLATE"].sub("PLATE-X", s)
        s = PATTERNS["PHONE"].sub("PHONE-X", s)
        s = PATTERNS["EMAIL"].sub("EMAIL-X", s)
        s = PATTERNS["NRIC_FIN"].sub("ID-X", s)
        s = PATTERNS["UNIT_NUMBER"].sub("#UNIT-X", s)
        return s

    def walk(o):
        if isinstance(o, dict):
            return {k: walk(v) for k, v in o.items()}
        if isinstance(o, list):
            return [walk(v) for v in o]
        if isinstance(o, str):
            return repl(o)
        return o
    return walk(out)


def fragment_scan_strings(strings) -> int:
    """Capitalized words right after or right before a code (ONT.11.4 items 4 and 6), e.g. 'DRIVER-NAME-1 Wei' or 'Ah Beng DRIVER-NAME-1'."""
    n = 0
    for s in strings:
        n += sum(1 for w in re.findall(r"\b" + CODE + r"\s+([A-Z][a-z]+)\b", s) if w not in NOT_NAMES)
        n += sum(1 for w in re.findall(r"\b([A-Z][a-z]+)\s+" + CODE + r"\b", s) if w not in NOT_NAMES)
    return n


def fragment_scan(case: dict) -> int:
    return fragment_scan_strings(_strings(case))


def dry_run(case: dict) -> dict:
    terms = dictionary(case)
    before = scan(case, terms)
    after_case = pseudonymize(case, terms)
    after = scan(after_case, terms)
    frag = fragment_scan(after_case)
    return {"dictionary_terms": len(terms), "before": before, "after": after, "fragments_after": frag,
            "clean": all(v == 0 for v in after.values()) and frag == 0}


def output_scan(strings, terms, allow_numbers=frozenset(), dictionary_only=()) -> dict:
    """Leak scan of generated outputs: dictionary terms, patterns and name fragments. Counts only.
    dictionary_only: texts checked for the case's real names and plate only (privacy.yaml documents the patterns with made-up examples)."""
    strings = [s for s in strings if isinstance(s, str)]
    counts = _count(strings, terms, allow_numbers)
    for k, v in _count(list(dictionary_only), terms, patterns=False).items():
        counts[k] += v
    counts["NAME_FRAGMENTS"] = fragment_scan_strings(strings)
    counts["total"] = sum(counts.values())
    return counts


def xlsx_strings(path) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    out.append(c.value)
                if c.comment is not None:
                    out.append(c.comment.text)
    return out


if __name__ == "__main__":
    import sys
    from kbload import KB
    kb = KB()
    print(json.dumps(dry_run(kb.disp002), indent=2))
    sys.exit(0)
